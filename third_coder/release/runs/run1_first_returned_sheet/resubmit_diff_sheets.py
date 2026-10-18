"""Cell-level diff of two returned coding sheets. Reads only the two sheets; opens no key file."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

OLD, NEW = sys.argv[1], sys.argv[2]
COLS = ['passage_id', 'passage_text', 'label', 'confidence', 'reason', 'lean']


def load(p):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    names = wb.sheetnames
    rows = list(wb['Coding'].iter_rows(values_only=True))
    return names, list(rows[0]), [r for r in rows[1:] if r[0] is not None]


def completeness(rows):
    n = ml = mlean = bad = 0
    for r in rows:
        n += 1
        lab = '' if r[2] is None else str(r[2]).strip().upper()
        lean = '' if r[5] is None else str(r[5]).strip().upper()
        if lab not in ('1', '0', 'U'):
            ml += 1
        if lab == 'U' and lean not in ('1', '0'):
            mlean += 1
        if r[3] is not None and str(r[3]).strip().capitalize() not in ('High', 'Medium', 'Low'):
            bad += 1
    return n, ml, mlean, bad


on, oh, orows = load(OLD)
nn, nh, nrows = load(NEW)
print('=== RESULTS ===')
print('sheetnames old', on, 'new', nn)
print('header same', oh == nh == COLS)
print('rows old', len(orows), 'new', len(nrows))
print('passage_id order same', [r[0] for r in orows] == [r[0] for r in nrows])
print('passage_text all same', all(a[1] == b[1] for a, b in zip(orows, nrows)))
print('completeness old rows=%d missing_label=%d U_without_lean=%d bad_conf=%d' % completeness(orows))
print('completeness new rows=%d missing_label=%d U_without_lean=%d bad_conf=%d' % completeness(nrows))
nchg = 0
for a, b in zip(orows, nrows):
    for i in range(2, 6):
        if a[i] != b[i]:
            nchg += 1
            print(f'CHANGE {a[0]} {COLS[i]}: {a[i]!r} -> {b[i]!r}')
print('changed cells', nchg)
print('=== END RESULTS ===')
