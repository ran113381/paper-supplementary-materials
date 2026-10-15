"""Selection script for the blind third-coder re-adjudication (executable form of J16_selection_rules.md).
Run:  D:\\python.exe -I selection_third_coder.py
Stops (non-zero exit) on any failed assertion or any situation the rules do not cover.
Never prints passage text and never prints the sample_id list of the non-concordant key or of the E1 key.
"""
import sys, os, re, csv, json, hashlib, copy, tempfile, datetime, platform
sys.stdout.reconfigure(encoding='utf-8')
import numpy as np
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

T0 = datetime.datetime.now().isoformat(timespec='seconds')
WB_PATH = r"E:\paper2  daxiu\73ae4696cf54747044a03a936c4bb958_744934475e1139b8c27898b82ae5bbf0_8\paper2_github_replication\data\genai_gold_standard_adjudication_workbook.xlsx"
WB_SHA = "aa072750c0a1ae9ac68dece24f060ef2c8e1d1bfb7a153244f9d69791620b0a0"
ROOT = r"E:\論文投稿  送審\paper2\R3_work"
OUT = os.path.join(ROOT, "third_coder")
SCRIPTS = os.path.join(ROOT, "third_coder_scripts")
RULES = os.path.join(ROOT, "J16_selection_rules.md")
SELF = os.path.abspath(__file__)

WL6 = ["passage_id", "passage_text", "label", "confidence", "reason", "lean"]
ENTRY = ["label", "confidence", "reason", "lean"]
BLOCK = ['hit', 'term', 'stratum', 'group', 'source', 'year', 'firm', 'model', 'llm', 'finbert', 'hybrid', 'pred',
         'coder', 'draft', 'final', 'adjud', 'consensus', 'gold', 'dictionary', 'matched', 'sample', 'metadata',
         'note', 'review', 'author', 'cell', 'nonhigh', 'reference', 'role', 'layer', 'subset', 'e1']
LEAK_RE = re.compile(r'GS_\d{4}|P2_(NON)?HIT_\d+')
CELLS = ['HIT-1', 'HIT-0', 'NM-0', 'RND-0']
CAL_TOT = {'HIT-1': 40, 'HIT-0': 30, 'NM-0': 20, 'RND-0': 10}
NH_TGT = {'HIT-1': 12, 'HIT-0': 20, 'NM-0': 12, 'RND-0': 6}
PRAC_Q = {'HIT-1': 8, 'HIT-0': 12, 'NM-0': 0, 'RND-0': 0}
STRAT = {'dictionary_hit_random': 'HIT', 'dictionary_non_hit_near_miss': 'NM', 'dictionary_non_hit_random': 'RND'}
FIXED_DT = datetime.datetime(2026, 10, 15, 0, 0, 0)

ASSERTS = {}   # id -> bool
CONTROLS = {}


def fail(msg):
    raise AssertionError(msg)


def chk(aid, cond, msg=''):
    if not cond:
        ASSERTS[aid] = False
        fail(f'{aid} FAILED {msg}')
    ASSERTS[aid] = True


def sha_file(p):
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


def norm_label(v):
    return str(v).strip().upper()


def norm_conf(v):
    return str(v).strip()


# ------------------------------------------------------------------ input
def read_sheet(wb, name, cols):
    ws = wb[name]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    idx = {}
    for c in cols:
        if c not in h:
            print(f'missing header {c!r} in sheet {name}; headers = {h}')
            fail(f'missing header {c} in {name}')
        idx[c] = h.index(c)
    rows = []
    for r in it:
        if all(v is None for v in r):
            continue
        rows.append({c: r[idx[c]] for c in cols})
    return rows


def load_input():
    wb = openpyxl.load_workbook(WB_PATH, read_only=True, data_only=True)
    ab = read_sheet(wb, 'Annotation_Blind', ['sample_id', 'mdna_passage', 'coder1_label', 'coder2_label',
                                              'coder1_confidence', 'coder2_confidence'])
    sm = read_sheet(wb, 'Sampling_Metadata', ['sample_id', 'sample_group'])
    cb = [(r[0], r[1]) for r in wb['Codebook'].iter_rows(min_row=2, max_row=11, max_col=2, values_only=True)]
    return ab, sm, cb


# ------------------------------------------------------------------ checks reusable by controls
def chk_A1(actual, expected):
    if actual != expected:
        fail('A1 hash mismatch')


def chk_A10(filelist):
    for n in ('analysis_third_coder.py', 'q7_construct_rescore.py'):
        if n not in filelist or filelist[n] <= 0:
            fail(f'A10 missing or empty {n}')


def check_no_leak(st, whitelist=WL6):
    h = st['header']
    if sorted(h) != sorted(whitelist) or len(set(h)) != len(h):
        fail(f'B5(a) header mismatch {h}')
    _leak_common(st)
    for c in ENTRY:
        i = h.index(c)
        if any(r[i] is not None and str(r[i]) != '' for r in st['rows']):
            fail('B5(d) entry column not empty')


def _leak_common(st):
    for x in st['header']:
        low = str(x).lower()
        for t in BLOCK:
            if t in low:
                fail(f'B5(b) header {x!r} contains {t!r}')
    for x in list(st['header']) + [v for r in st['rows'] for v in r]:
        if isinstance(x, str) and LEAK_RE.search(x):
            fail('B5(c) identifier pattern in a cell')
    if st['comments'] or st['defined_names'] or st['hidden']:
        fail('B5(e) comments / defined names / hidden content')
    p = st['props']
    if p['creator'] != 'coordinator':
        fail('B5(f) creator')
    for k in ('lastModifiedBy', 'title', 'subject', 'keywords', 'description', 'category'):
        if p[k] not in (None, ''):
            fail(f'B5(f) property {k}')


def check_no_leak_answers(st):
    _leak_common(st)


def chk_B1(st, n, sheetname, header, idpat=None):
    if st['sheetnames'] != [sheetname]:
        fail('B1 sheet names')
    if len(st['rows']) != n:
        fail(f'B1 row count {len(st["rows"])} != {n}')
    if st['header'] != header:
        fail('B1 header')


def chk_B4(st, idmap, wbtext):
    for r in st['rows']:
        if r[1] != wbtext[idmap[r[0]]]:
            fail('B4 text mismatch')


def chk_B6_order(seq_rep, seq_cod):
    if set(seq_rep) != set(seq_cod):
        fail('B6 sets differ')
    if list(seq_rep) == list(seq_cod):
        fail('B6 order identical')


def chk_B8(nonc, cal, prac, e1):
    sets = [nonc, cal, prac, e1]
    for i in range(4):
        for j in range(i + 1, 4):
            if sets[i] & sets[j]:
                fail('B8 sets overlap')
    if (len(nonc), len(cal), len(prac), len(e1)) != (123, 100, 20, 214):
        fail('B8 sizes')
    if 'GS_0438' not in nonc:
        fail('B8 GS_0438')


def chk_B16(e1layer, cal, e1, prac, roles, cal_cells, ref, strat):
    cal_e = {s for s in cal if cal_cells[s] in ('NM-0', 'RND-0')}
    if len(e1layer) != 244 or len(cal_e) != 30 or len(e1) != 214:
        fail('B16 sizes')
    if e1layer != (cal_e | e1):
        fail('B16 coverage')
    if e1layer & prac:
        fail('B16 practice overlap')
    for s in e1layer:
        if ref[s] != 0 or strat[s] not in ('NM', 'RND') or roles[s] in ('unused', 'nonconcordant', 'practice'):
            fail('B16 layer member invalid')


# ------------------------------------------------------------------ selection (S1-S7)
def draw(pool_ids, k, rng):
    pool = sorted(pool_ids)
    if k > len(pool):
        fail(f'A6 pool {len(pool)} < k {k}')
    perm = rng.permutation(len(pool))
    return [pool[i] for i in perm[:k]]


def classify(ab, sm):
    D = {}
    for a_, m in zip(ab, sm):
        sid = a_['sample_id']
        c1 = norm_label(a_['coder1_label']); c2 = norm_label(a_['coder2_label'])
        if c1 not in ('0', '1', 'U') or c2 not in ('0', '1', 'U'):
            fail(f'bad label at {sid}')
        f1 = norm_conf(a_['coder1_confidence']); f2 = norm_conf(a_['coder2_confidence'])
        if f1 not in ('High', 'Medium', 'Low') or f2 not in ('High', 'Medium', 'Low'):
            fail(f'bad confidence at {sid}')
        if m['sample_group'] not in STRAT:
            fail(f'bad sample_group at {sid}')
        t = a_['mdna_passage']
        D[sid] = dict(text=t, c1=c1, c2=c2, nonhigh=(f1 != 'High' or f2 != 'High'), stratum=STRAT[m['sample_group']],
                      conc=(c1 == c2 and c1 in ('0', '1')))
        D[sid]['ref'] = int(c1) if D[sid]['conc'] else None
        D[sid]['cell'] = None
        if D[sid]['conc']:
            if D[sid]['stratum'] == 'HIT':
                D[sid]['cell'] = 'HIT-1' if D[sid]['ref'] == 1 else 'HIT-0'
            else:
                if D[sid]['ref'] == 1:
                    fail('design assumption violated: concordant positive in non-matched stratum')
                D[sid]['cell'] = D[sid]['stratum'] + '-0'
    # dup rule
    conc_ids = sorted(s for s in D if D[s]['conc'])
    nonc_texts = {D[s]['text'] for s in D if not D[s]['conc']}
    groups = {}
    for s in conc_ids:
        groups.setdefault(D[s]['text'], []).append(s)
    for s in D:
        D[s]['dup'] = False
    for t, g in groups.items():
        if len(g) >= 2:
            for s in sorted(g)[1:]:
                D[s]['dup'] = True
    for s in conc_ids:
        if D[s]['text'] in nonc_texts:
            D[s]['dup'] = True
    for s in D:
        D[s]['e1layer'] = bool(D[s]['conc'] and D[s]['stratum'] in ('NM', 'RND'))
    return D


def select(D, seed_sel, seed_rep, cb):
    rng = np.random.default_rng(seed_sel)
    role = {s: None for s in D}
    for s in D:
        if not D[s]['conc']:
            role[s] = 'nonconcordant'
    short = []
    knh = {}

    def pool(cell, nh):
        return [s for s in D if D[s]['conc'] and D[s]['cell'] == cell and not D[s]['dup'] and role[s] is None
                and D[s]['nonhigh'] == nh]
    # S1
    for cell in CELLS:
        nh = pool(cell, True); hh = pool(cell, False)
        k_nh = min(NH_TGT[cell], len(nh))
        A = draw(nh, k_nh, rng)
        k_hh = CAL_TOT[cell] - k_nh
        B = draw(hh, k_hh, rng)
        if k_nh < NH_TGT[cell]:
            short.append(f'SHORTFALL S1 {cell} nonhigh {k_nh}/{NH_TGT[cell]}')
        knh[cell] = k_nh
        for s in A + B:
            role[s] = 'calibration'
    # S2
    s2short = {}
    for cell in ('HIT-1', 'HIT-0'):
        hh = pool(cell, False)
        q = PRAC_Q[cell]
        if len(hh) >= q:
            P = draw(hh, q, rng)
        else:
            nhr = pool(cell, True)
            P = draw(hh, len(hh), rng) + draw(nhr, q - len(hh), rng)
            short.append(f'SHORTFALL S2 {cell} bothhigh {len(hh)}/{q}')
            s2short[cell] = True
        for s in P:
            role[s] = 'practice'
    ASSERTS['A6'] = True
    # S3
    e1_set = [s for s in D if D[s]['e1layer'] and role[s] is None]
    cal_set = {s for s in D if role[s] == 'calibration'}
    prac_set = {s for s in D if role[s] == 'practice'}
    if not (len(e1_set) == 214 and len({s for s in cal_set if D[s]['e1layer']}) == 30
            and len({s for s in prac_set if D[s]['e1layer']}) == 0):
        fail('S3 assertion')
    for s in e1_set:
        role[s] = 'e1'
    for s in D:
        if role[s] is None:
            role[s] = 'unused'
    # S4
    rows = sorted(s for s in D if role[s] in ('nonconcordant', 'calibration', 'e1'))
    if len(rows) != 437:
        fail('S4 rows != 437')
    perm = rng.permutation(437)
    ordered = [rows[i] for i in perm]
    # S5 (separate generator)
    rng_r = np.random.default_rng(seed_rep)
    perm_r = rng_r.permutation(437)
    ordered_r = [rows[i] for i in perm_r]
    # S6
    prows = sorted(s for s in D if role[s] == 'practice')
    perm_p = rng.permutation(20)
    pordered = [prows[i] for i in perm_p]
    return dict(role=role, short=short, knh=knh, s2short=s2short, ordered=ordered, ordered_r=ordered_r,
                pordered=pordered)


# ------------------------------------------------------------------ writing
def new_wb(sheetname):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheetname
    wb.properties.creator = 'coordinator'
    for k in ('lastModifiedBy', 'title', 'subject', 'keywords', 'description', 'category'):
        setattr(wb.properties, k, None)
    wb.properties.created = FIXED_DT
    wb.properties.modified = FIXED_DT
    return wb, ws


def put(ws, r, c, v, sid=None):
    cell = ws.cell(row=r, column=c)
    try:
        cell.value = v
    except Exception as e:
        raise AssertionError(f'cannot write cell for {sid}: {type(e).__name__}') from None
    if isinstance(v, str) and v.startswith('='):
        cell.data_type = 's'


def write_entry_sheet(path, sheetname, ids, texts, sids):
    wb, ws = new_wb(sheetname)
    for j, h in enumerate(WL6, 1):
        put(ws, 1, j, h)
    for i, (pid, t, sid) in enumerate(zip(ids, texts, sids), 2):
        put(ws, i, 1, pid, sid)
        put(ws, i, 2, t, sid)
    n = len(ids)
    for col, lst in (('C', '"1,0,U"'), ('D', '"High,Medium,Low"'), ('F', '"1,0"')):
        dv = DataValidation(type='list', formula1=lst, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{n + 1}')
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 100
    wb.save(path)


def write_simple(path, sheetname, header, rows):
    wb, ws = new_wb(sheetname)
    for j, h in enumerate(header, 1):
        put(ws, 1, j, h)
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            put(ws, i, j, v)
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 100
    wb.save(path)


def write_csv(path, header, rows):
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(header)
        w.writerows(rows)


def cb_text(cb, ref):
    row = cb[2] if ref == 1 else cb[3]          # CB03 / CB04
    return f'{row[0]} | {row[1]}'


def build(root, D, S, cb):
    for sub in ('to_coder', 'sealed', os.path.join('sealed', 'replacement'), 'manifest'):
        os.makedirs(os.path.join(root, sub), exist_ok=True)
    role = S['role']
    tc = os.path.join(root, 'to_coder'); se = os.path.join(root, 'sealed'); rp = os.path.join(se, 'replacement')
    # S7 codebook
    rows = [('Gold-standard annotation codebook', None)] + [(str(a_), str(b_)) for a_, b_ in cb[0:10]]
    write_simple(os.path.join(tc, 'codebook_for_coder.xlsx'), 'Codebook', ['item', 'text'], rows)
    # coding sheets
    for fname, order, pre, keysuffix, folder in (('coding_sheet.xlsx', S['ordered'], 'A', '', tc),
                                                  ('coding_sheet_replacement.xlsx', S['ordered_r'], 'R', '_replacement', rp)):
        ids = [f'{pre}{j + 1:03d}' for j in range(437)]
        write_entry_sheet(os.path.join(folder, fname), 'Coding', ids, [D[s]['text'] for s in order], order)
        idmap = dict(zip(order, ids))
        kfolder = rp if pre == 'R' else se
        by = sorted(order, key=lambda s: idmap[s])
        write_csv(os.path.join(kfolder, f'key_calibration{keysuffix}.csv'),
                  ['passage_id', 'sample_id', 'cell', 'reference_label', 'nonhigh', 'e1_layer'],
                  [[idmap[s], s, D[s]['cell'], D[s]['ref'], D[s]['nonhigh'], D[s]['e1layer']] for s in by if role[s] == 'calibration'])
        write_csv(os.path.join(kfolder, f'key_nonconcordant{keysuffix}.csv'), ['passage_id', 'sample_id'],
                  [[idmap[s], s] for s in by if role[s] == 'nonconcordant'])
        write_csv(os.path.join(kfolder, f'key_e1{keysuffix}.csv'), ['passage_id', 'sample_id', 'stratum', 'reference_label'],
                  [[idmap[s], s, D[s]['stratum'], 0] for s in by if role[s] == 'e1'])
    # practice
    pids = [f'T{j + 1:02d}' for j in range(20)]
    po = S['pordered']
    write_entry_sheet(os.path.join(tc, 'practice_sheet.xlsx'), 'Practice', pids, [D[s]['text'] for s in po], po)
    write_simple(os.path.join(tc, 'practice_answers.xlsx'), 'Answers', ['passage_id', 'feedback_label', 'codebook_text'],
                 [[pid, D[s]['ref'], cb_text(cb, D[s]['ref'])] for pid, s in zip(pids, po)])
    write_csv(os.path.join(se, 'key_practice.csv'),
              ['passage_id', 'sample_id', 'cell', 'reference_label', 'codebook_row', 'codebook_text'],
              [[pid, s, D[s]['cell'], D[s]['ref'], 'CB03' if D[s]['ref'] == 1 else 'CB04', cb_text(cb, D[s]['ref'])]
               for pid, s in zip(pids, po)])
    write_csv(os.path.join(se, 'selection_register.csv'),
              ['sample_id', 'role', 'cell', 'nonhigh', 'dup_ineligible', 'e1_layer'],
              [[s, role[s], D[s]['cell'] or '', D[s]['nonhigh'], D[s]['dup'], D[s]['e1layer']] for s in sorted(D)])


# ------------------------------------------------------------------ reading back
def load_struct(path):
    wb = openpyxl.load_workbook(path)
    names = wb.sheetnames
    hidden = False; comments = False
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible':
            hidden = True
        if any(d.hidden for d in ws.row_dimensions.values()) or any(d.hidden for d in ws.column_dimensions.values()):
            hidden = True
        for row in ws.iter_rows():
            for c in row:
                if c.comment is not None:
                    comments = True
    ws0 = wb.worksheets[0]
    rows = [list(r) for r in ws0.iter_rows(values_only=True)]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    p = wb.properties
    dn = list(wb.defined_names.keys()) if hasattr(wb.defined_names, 'keys') else list(wb.defined_names)
    return dict(sheetnames=names, header=rows[0], rows=rows[1:], comments=comments, defined_names=dn, hidden=hidden,
                props=dict(creator=p.creator, lastModifiedBy=p.lastModifiedBy, title=p.title, subject=p.subject,
                           keywords=p.keywords, description=p.description, category=p.category))


def canon_hash(st):
    lines = ['\t'.join('' if v is None else str(v) for v in r) for r in [st['header']] + st['rows']]
    return hashlib.sha256(('\n'.join(lines) + '\n').encode('utf-8')).hexdigest()


def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        rd = csv.DictReader(f)
        return rd.fieldnames, list(rd)


def tree_hashes(root):
    out = {}
    for dp, _, fns in os.walk(root):
        for fn in fns:
            p = os.path.join(dp, fn)
            rel = os.path.relpath(p, root)
            if rel.startswith('manifest') or fn.startswith('seed_'):
                continue
            out[rel] = (sha_file(p), canon_hash(load_struct(p)) if fn.endswith('.xlsx') else None)
    return out


# ------------------------------------------------------------------ main
def main():
    # seeds
    se = os.path.join(OUT, 'sealed')
    seedp = {k: os.path.join(se, f'seed_{k}.txt') for k in ('selection', 'replacement', 'analysis')}
    seeds = {}
    for k, p in seedp.items():
        if not os.path.exists(p):
            fail(f'missing seed file {p}')
        seeds[k] = int(open(p, encoding='utf-8').read().strip())
    # A1
    chk_A1(sha_file(WB_PATH), WB_SHA)
    ASSERTS['A1'] = True
    ab, sm, cb = load_input()
    # A2
    sids = [r['sample_id'] for r in ab]
    chk('A2', len(ab) == 600 and len(sm) == 600 and sids == [r['sample_id'] for r in sm] and len(set(sids)) == 600
        and all(re.fullmatch(r'GS_\d{4}', str(s)) for s in sids))
    D = classify(ab, sm)
    nconc = sum(D[s]['conc'] for s in D)
    chk('A3', nconc == 477 and len(D) - nconc == 123 and len(D) == 600)
    chk('A4', (not D['GS_0438']['conc']) and D['GS_0438']['c1'] == 'U' and D['GS_0438']['c2'] == 'U')
    ASSERTS['A5'] = True   # enforced inside classify() (stops if violated)
    chk('A7', all(isinstance(D[s]['text'], str) and 0 < len(D[s]['text']) <= 32767 for s in D))
    e1l = [s for s in D if D[s]['e1layer']]
    chk('A8', len(e1l) == 244, len(e1l))
    chk('A9', not any(D[s]['dup'] for s in e1l))
    flist = {n: os.path.getsize(os.path.join(SCRIPTS, n)) for n in os.listdir(SCRIPTS)} if os.path.isdir(SCRIPTS) else {}
    try:
        chk_A10(flist)
        ASSERTS['A10'] = True
    except AssertionError:
        ASSERTS['A10'] = False
        raise
    # nothing but seeds may exist yet
    for dp, _, fns in os.walk(OUT):
        for fn in fns:
            if not fn.startswith('seed_'):
                fail(f'output directory already holds {os.path.join(dp, fn)}')
    # selection + write
    S = select(D, seeds['selection'], seeds['replacement'], cb)
    build(OUT, D, S, cb)
    role = S['role']
    wbtext = {s: D[s]['text'] for s in D}
    tc = os.path.join(OUT, 'to_coder'); rp = os.path.join(se, 'replacement')

    # ------------- post-write assertions on files from disk
    def keys(folder, suffix):
        out = {}
        for n, cols in (('key_calibration', ['passage_id', 'sample_id', 'cell', 'reference_label', 'nonhigh', 'e1_layer']),
                        ('key_nonconcordant', ['passage_id', 'sample_id']),
                        ('key_e1', ['passage_id', 'sample_id', 'stratum', 'reference_label'])):
            h, r = read_csv(os.path.join(folder, f'{n}{suffix}.csv'))
            if h != cols:
                fail(f'{n} header')
            out[n] = r
        return out

    def check_sheet(path, folder, suffix, pre):
        st = load_struct(path)
        chk_B1(st, 437, 'Coding', WL6)
        ids = [r[0] for r in st['rows']]
        if len(set(ids)) != 437 or ids != [f'{pre}{j + 1:03d}' for j in range(437)]:
            fail('B2 ids')
        K = keys(folder, suffix)
        n_, c_, e_ = K['key_nonconcordant'], K['key_calibration'], K['key_e1']
        if (len(n_), len(c_), len(e_)) != (123, 100, 214):
            fail('B3 key sizes')
        pn, pc, pe = ({r['passage_id'] for r in x} for x in (n_, c_, e_))
        if pn & pc or pn & pe or pc & pe or (pn | pc | pe) != set(ids):
            fail('B3 passage_id sets')
        idmap = {r['passage_id']: r['sample_id'] for r in n_ + c_ + e_}
        sidset = {idmap[i] for i in ids}
        reg_h, reg = read_csv(os.path.join(se, 'selection_register.csv'))
        rr = {r['sample_id']: r['role'] for r in reg}
        want = {s for s in rr if rr[s] in ('nonconcordant', 'calibration', 'e1')}
        if sidset != want or {r['sample_id'] for r in n_} != {s for s in rr if rr[s] == 'nonconcordant'} \
                or {r['sample_id'] for r in c_} != {s for s in rr if rr[s] == 'calibration'} \
                or {r['sample_id'] for r in e_} != {s for s in rr if rr[s] == 'e1'}:
            fail('B3 sample_id sets')
        chk_B4(st, idmap, wbtext)
        check_no_leak(st)
        return st, idmap, [idmap[i] for i in ids], K

    st1, idmap1, seq1, K1 = check_sheet(os.path.join(tc, 'coding_sheet.xlsx'), se, '', 'A')
    ASSERTS['B1'] = ASSERTS['B2'] = ASSERTS['B3'] = ASSERTS['B4'] = ASSERTS['B5'] = True
    st2, idmap2, seq2, K2 = check_sheet(os.path.join(rp, 'coding_sheet_replacement.xlsx'), rp, '_replacement', 'R')
    chk_B6_order(seq2, seq1)
    # replacement keys map to the same sets as S4 keys
    for n in ('key_calibration', 'key_nonconcordant', 'key_e1'):
        if {r['sample_id'] for r in K1[n]} != {r['sample_id'] for r in K2[n]}:
            fail('B6 replacement key sets')
    ASSERTS['B6'] = True
    # practice
    stp = load_struct(os.path.join(tc, 'practice_sheet.xlsx'))
    chk_B1(stp, 20, 'Practice', WL6)
    pids = [r[0] for r in stp['rows']]
    if pids != [f'T{j + 1:02d}' for j in range(20)]:
        fail('B7 ids')
    check_no_leak(stp)
    kph, kp = read_csv(os.path.join(se, 'key_practice.csv'))
    if kph != ['passage_id', 'sample_id', 'cell', 'reference_label', 'codebook_row', 'codebook_text']:
        fail('key_practice header')
    pmap = {r['passage_id']: r for r in kp}
    chk_B4(stp, {k: v['sample_id'] for k, v in pmap.items()}, wbtext)
    if any(D[r['sample_id']]['stratum'] != 'HIT' for r in kp):
        fail('B7 stratum')
    pc = {c: sum(r['cell'] == c for r in kp) for c in CELLS}
    if pc != {'HIT-1': 8, 'HIT-0': 12, 'NM-0': 0, 'RND-0': 0}:
        fail('B7 cell counts')
    ASSERTS['B7'] = True
    # register-based sets
    rh, reg = read_csv(os.path.join(se, 'selection_register.csv'))
    sets = {k: {r['sample_id'] for r in reg if r['role'] == k} for k in ('nonconcordant', 'calibration', 'practice', 'e1', 'unused')}
    chk_B8(sets['nonconcordant'], sets['calibration'], sets['practice'], sets['e1'])
    if 'GS_0438' not in sets['nonconcordant'] or any('GS_0438' in sets[k] for k in ('calibration', 'practice', 'e1', 'unused')):
        fail('B8 GS_0438')
    ASSERTS['B8'] = True
    # B9
    cal_cells = {r['sample_id']: r['cell'] for r in K1['key_calibration']}
    for c in CELLS:
        n_c = sum(v == c for v in cal_cells.values())
        n_nh = sum(1 for r in K1['key_calibration'] if r['cell'] == c and r['nonhigh'] == 'True')
        if n_c != CAL_TOT[c] or n_nh != S['knh'][c]:
            fail(f'B9 {c}')
    ASSERTS['B9'] = True
    # B10
    for r in kp:
        if D[r['sample_id']]['nonhigh'] and r['cell'] not in S['s2short']:
            fail('B10')
    ASSERTS['B10'] = True
    # B11
    drawn = sets['calibration'] | sets['practice'] | sets['e1']
    if any(D[s]['dup'] for s in drawn):
        fail('B11 dup in drawn set')
    ctexts = [r[1] for r in st1['rows']]; ptexts = [r[1] for r in stp['rows']]
    if len(set(ctexts)) != 437 or len(set(ptexts)) != 20 or set(ptexts) & set(ctexts):
        fail('B11 texts')
    ASSERTS['B11'] = True
    # B13
    stc = load_struct(os.path.join(tc, 'codebook_for_coder.xlsx'))
    if stc['sheetnames'] != ['Codebook'] or stc['header'] != ['item', 'text'] or len(stc['rows']) != 11:
        fail('B13 shape')
    if stc['rows'][0][0] != 'Gold-standard annotation codebook' or stc['rows'][0][1] is not None:
        fail('B13 title row')
    for i in range(10):
        if stc['rows'][i + 1] != [str(cb[i][0]), str(cb[i][1])]:
            fail(f'B13 row {i + 3}')
    p = stc['props']
    if p['creator'] != 'coordinator' or any(p[k] not in (None, '') for k in ('lastModifiedBy', 'title', 'subject', 'keywords', 'description', 'category')):
        fail('B13 props')
    ASSERTS['B13'] = True
    # B14
    for seq in (seq1, seq2, [pmap[i]['sample_id'] for i in pids]):
        if seq == sorted(seq):
            fail('B14 identity order')
    ASSERTS['B14'] = True
    # B15
    cnt = {k: len(v) for k, v in sets.items()}
    if len(reg) != 600 or cnt != {'nonconcordant': 123, 'calibration': 100, 'practice': 20, 'e1': 214, 'unused': 143} \
            or sum(r['e1_layer'] == 'True' for r in reg) != 244:
        fail(f'B15 {cnt}')
    ASSERTS['B15'] = True
    # B16
    rroles = {r['sample_id']: r['role'] for r in reg}
    e1layer = {r['sample_id'] for r in reg if r['e1_layer'] == 'True'}
    chk_B16(e1layer, sets['calibration'], sets['e1'], sets['practice'], rroles,
            {r['sample_id']: r['cell'] for r in reg}, {s: (D[s]['ref'] if D[s]['conc'] else None) for s in D},
            {s: D[s]['stratum'] for s in D})
    ASSERTS['B16'] = True
    # B17
    sta = load_struct(os.path.join(tc, 'practice_answers.xlsx'))
    chk_B1(sta, 20, 'Answers', ['passage_id', 'feedback_label', 'codebook_text'])
    if [r[0] for r in sta['rows']] != pids:
        fail('B17 ids')
    for r in sta['rows']:
        k = pmap[r[0]]
        if str(r[1]) != k['reference_label'] or r[2] != k['codebook_text']:
            fail('B17 content')
    check_no_leak_answers(sta)
    ASSERTS['B17'] = True

    # ------------- B12 determinism (second run in a temporary directory, removed automatically)
    h1 = tree_hashes(OUT)
    with tempfile.TemporaryDirectory() as td:
        S2 = select(D, seeds['selection'], seeds['replacement'], cb)
        build(td, D, S2, cb)
        h2 = tree_hashes(td)
    keyset = set(h1)
    if keyset != set(h2):
        fail('B12 file sets differ')
    for k in keyset:
        if k.endswith('.xlsx'):
            if h1[k][1] != h2[k][1]:
                fail(f'B12 canonical hash differs {k}')
        elif h1[k][0] != h2[k][0]:
            fail(f'B12 byte hash differs {k}')
    ASSERTS['B12'] = True

    # ------------- should-fail controls (in-memory copies)
    def must_raise(cid, fn):
        try:
            fn()
        except AssertionError:
            CONTROLS[cid] = True
            return
        CONTROLS[cid] = False
        raise SystemExit(f'control {cid} did not raise')

    rows_all = [list(r) for r in st1['rows']]

    # C1: joined table
    joined = dict(sheetnames=['x'], header=list(ab[0].keys()) + [k for k in sm[0].keys() if k != 'sample_id'],
                  rows=[[a_[k] for k in ab[0].keys()] + [m[k] for k in sm[0].keys() if k != 'sample_id'] for a_, m in zip(ab, sm)],
                  comments=False, defined_names=[], hidden=False, props=copy.deepcopy(st1['props']))
    must_raise('C1', lambda: check_no_leak(joined))
    t2 = copy.deepcopy(st1); t2['header'] = t2['header'] + ['dictionary_hit']; t2['rows'] = [r + ['x'] for r in t2['rows']]
    must_raise('C2', lambda: check_no_leak(t2))
    t3 = copy.deepcopy(st1); t3['rows'][0][0] = idmap1[t3['rows'][0][0]]
    must_raise('C3', lambda: check_no_leak(t3))
    t4 = copy.deepcopy(st1); t4['rows'] = t4['rows'][:436]
    must_raise('C4', lambda: chk_B1(t4, 437, 'Coding', WL6))
    cal_one = sorted(sets['calibration'])[0]
    must_raise('C5', lambda: chk_B8(sets['nonconcordant'], sets['calibration'], sets['practice'] | {cal_one}, sets['e1']))
    must_raise('C6', lambda: chk_A1(sha_file(WB_PATH), WB_SHA[:-1] + ('0' if WB_SHA[-1] != '0' else '1')))
    t7 = copy.deepcopy(st1); t7['rows'][0][1] = t7['rows'][0][1] + 'X'
    must_raise('C7', lambda: chk_B4(t7, idmap1, wbtext))
    e1_minus = set(sorted(sets['e1'])[1:])
    must_raise('C8', lambda: chk_B16(e1layer, sets['calibration'], e1_minus, sets['practice'], rroles,
                                      {r['sample_id']: r['cell'] for r in reg}, {s: (D[s]['ref'] if D[s]['conc'] else None) for s in D},
                                      {s: D[s]['stratum'] for s in D}))
    must_raise('C9', lambda: chk_B6_order(seq1, seq1))
    t10 = copy.deepcopy(st1); t10['header'] = t10['header'] + ['e1_layer']; t10['rows'] = [r + [None] for r in t10['rows']]
    must_raise('C10', lambda: check_no_leak(t10))
    t11 = copy.deepcopy(sta); t11['header'] = ['reference_label' if h == 'feedback_label' else h for h in t11['header']]
    must_raise('C11', lambda: check_no_leak_answers(t11))
    t12 = copy.deepcopy(sta); t12['rows'][0][0] = pmap[t12['rows'][0][0]]['sample_id']
    must_raise('C12', lambda: check_no_leak_answers(t12))
    fl2 = dict(flist); fl2.pop('analysis_third_coder.py')
    must_raise('C14', lambda: chk_A10(fl2))

    # ------------- manifest + RESULTS
    outfiles = {}
    for dp, _, fns in os.walk(OUT):
        for fn in sorted(fns):
            p = os.path.join(dp, fn)
            rel = os.path.relpath(p, OUT).replace('\\', '/')
            if rel.startswith('manifest') or fn.startswith('seed_'):
                continue
            outfiles[rel] = sha_file(p)
    canon = {rel: canon_hash(load_struct(os.path.join(OUT, rel))) for rel in outfiles if rel.endswith('.xlsx')}
    seed_sha = {k: sha_file(p) for k, p in seedp.items()}
    pcell = {c: sum(1 for s in D if role[s] == 'practice' and D[s]['cell'] == c) for c in CELLS}
    ccell = {c: sum(1 for s in D if role[s] == 'calibration' and D[s]['cell'] == c) for c in CELLS}
    e1cnt = {st: sum(1 for s in e1layer if D[s]['stratum'] == st) for st in ('NM', 'RND')}
    e1role = {st: {r: sum(1 for s in e1layer if D[s]['stratum'] == st and role[s] == r) for r in ('calibration', 'e1')} for st in ('NM', 'RND')}
    ndup = sum(D[s]['dup'] for s in D)
    nP = sum(1 for k in ASSERTS if ASSERTS[k]); nF = sum(1 for k in ASSERTS if not ASSERTS[k])
    cP = sum(1 for k in CONTROLS if CONTROLS[k]); cF = sum(1 for k in CONTROLS if not CONTROLS[k])
    assert len(ASSERTS) == 27 and len(CONTROLS) == 13, (len(ASSERTS), len(CONTROLS))
    L = ['=== RESULTS ===',
         f'workbook_sha256={WB_SHA}  seed_selection_sha256={seed_sha["selection"]}  seed_replacement_sha256={seed_sha["replacement"]}  seed_analysis_sha256={seed_sha["analysis"]}',
         f'nonconcordant={len(D) - nconc} concordant={nconc} GS_0438={role["GS_0438"]}',
         'calibration per cell: ' + ' | '.join(f'{c} {ccell[c]} (nonhigh {S["knh"][c]})' for c in CELLS),
         'practice per cell: ' + ' | '.join(f'{c} {pcell[c]}' for c in CELLS),
         f'e1_layer={len(e1layer)} (NM {e1cnt["NM"]}, RND {e1cnt["RND"]}) = calibration {len(e1layer & sets["calibration"])} + e1 {len(e1layer & sets["e1"])} (practice {len(e1layer & sets["practice"])})',
         f'coding sheet rows={len(st1["rows"])}  replacement coding sheet rows={len(st2["rows"])}  practice rows={len(stp["rows"])}  answer sheet rows={len(sta["rows"])}',
         'role counts: ' + ' | '.join(f'{k} {cnt[k]}' for k in ('nonconcordant', 'calibration', 'practice', 'e1', 'unused')),
         'SHORTFALL lines: ' + ('; '.join(S['short']) if S['short'] else 'none'),
         f'dup_ineligible: {ndup} (in e1_layer: {sum(D[s]["dup"] for s in e1layer)})',
         f'assertions A1-A10, B1-B17: {nP} PASS / {nF} FAIL',
         f'controls C1-C12, C14: {cP} PASS / {cF} FAIL',
         'sha256 per output file:'] + [f'{k}={v}' for k, v in sorted(outfiles.items())] + \
        ['canonical content hash per xlsx:'] + [f'{k}={v}' for k, v in sorted(canon.items())] + \
        [f'python={platform.python_version()} numpy={np.__version__} openpyxl={openpyxl.__version__}', '=== END RESULTS ===']
    block = '\n'.join(L)
    man = dict(run_start=T0, workbook_sha256=WB_SHA, seed_files_sha256=seed_sha,
               rules_file_sha256=sha_file(RULES), selection_script_sha256=sha_file(SELF),
               analysis_script_sha256=sha_file(os.path.join(SCRIPTS, 'analysis_third_coder.py')),
               q7_script_sha256=sha_file(os.path.join(SCRIPTS, 'q7_construct_rescore.py')),
               output_files_sha256=outfiles, canonical_content_hash_xlsx=canon,
               counts=dict(calibration_per_cell=ccell, calibration_nonhigh_per_cell=S['knh'], practice_per_cell=pcell,
                           role_counts=cnt, e1_layer_by_stratum=e1cnt, e1_by_stratum_and_role=e1role,
                           dup_ineligible=ndup),
               shortfall_lines=S['short'],
               versions=dict(python=platform.python_version(), numpy=np.__version__, openpyxl=openpyxl.__version__),
               results_block=block)
    with open(os.path.join(OUT, 'manifest', 'selection_manifest.json'), 'w', encoding='utf-8') as f:
        json.dump(man, f, ensure_ascii=False, indent=2, sort_keys=True)
    print(block)


main()
