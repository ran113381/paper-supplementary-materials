"""Analysis script for the blind third-coder re-adjudication (protocol J16, sections 6.5, 7, 8.2(vi)(viii)(ix)).

Run with:  D:\\python.exe -I analysis_third_coder.py --stage 1|2 ...

Stage 1 opens ONLY key_calibration.csv (calibration rows of the returned sheet), computes the
qualification statistics and writes a qualification record whose SHA-256 is appended to a log.
Stage 2 opens key_nonconcordant.csv and key_e1.csv ONLY if the record exists and its hash matches the log.
Stage 2 runs whatever the Stage 1 outcome; the outcome decides only whether final_labels.csv is written.
Read-only on every key / sheet / workbook file; writes only into --outdir.
The model comparison (protocol 8.5) is done separately by q7_construct_rescore.py via --gold-file.
"""
import sys, os, re, csv, json, math, hashlib, argparse, datetime
sys.stdout.reconfigure(encoding='utf-8')
import numpy as np
import openpyxl

THRESH = 0.667
Z = 1.959963984540054
ap = argparse.ArgumentParser()
ap.add_argument('--stage', type=int, required=True, choices=[1, 2])
ap.add_argument('--sheet', required=True, help='returned coding sheet (xlsx, sheet Coding)')
ap.add_argument('--keydir', required=True, help='sealed directory (or sealed\\replacement with --replacement)')
ap.add_argument('--replacement', action='store_true', help='use *_replacement.csv keys (protocol 10.5)')
ap.add_argument('--workbook', required=True, help='gold-standard workbook (read-only; coder1/coder2 labels)')
ap.add_argument('--seed-file', required=True, help='seed_analysis.txt')
ap.add_argument('--outdir', required=True)
a = ap.parse_args()

os.makedirs(a.outdir, exist_ok=True)
suffix = '_replacement' if a.replacement else ''
REC = os.path.join(a.outdir, 'qualification_record.json')
LOG = os.path.join(a.outdir, 'qualification_record_hash_log.txt')


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


def norm_label(v):
    return '' if v is None else str(v).strip().upper()


def norm_conf(v):
    return '' if v is None else str(v).strip().capitalize()


def read_key(name, cols):
    p = os.path.join(a.keydir, name.replace('.csv', suffix + '.csv'))
    with open(p, encoding='utf-8-sig', newline='') as f:
        rd = csv.DictReader(f)
        assert rd.fieldnames == cols, (p, rd.fieldnames)
        return list(rd)


def read_sheet(ids=None):
    """returns dict passage_id -> dict(label, confidence, reason, lean); only ids if given."""
    wb = openpyxl.load_workbook(a.sheet, read_only=True, data_only=True)
    assert wb.sheetnames == ['Coding'], wb.sheetnames
    it = wb['Coding'].iter_rows(values_only=True)
    h = list(next(it))
    assert h == ['passage_id', 'passage_text', 'label', 'confidence', 'reason', 'lean'], h
    out = {}
    for r in it:
        pid = r[0]
        if pid is None:
            continue
        if ids is not None and pid not in ids:
            continue
        assert pid not in out, f'duplicate passage_id {pid}'
        out[pid] = dict(label=norm_label(r[2]), confidence=norm_conf(r[3]),
                        reason='' if r[4] is None else str(r[4]).strip(), lean=norm_label(r[5]))
    return out


def sheet_completeness():
    """counts only; no label values are returned"""
    wb = openpyxl.load_workbook(a.sheet, read_only=True, data_only=True)
    n = miss_label = miss_lean = bad = 0
    for r in list(wb['Coding'].iter_rows(values_only=True))[1:]:
        if r[0] is None:
            continue
        n += 1
        lab = norm_label(r[2]); lean = norm_label(r[5])
        if lab not in ('1', '0', 'U'):
            miss_label += 1
        if lab == 'U' and lean not in ('1', '0'):
            miss_lean += 1
        if r[3] is not None and norm_conf(r[3]) not in ('High', 'Medium', 'Low'):
            bad += 1
    return n, miss_label, miss_lean, bad


# ---------------- statistics ----------------
def kripp_alpha_nominal(pairs):
    """pairs: list of (a, b) with None = missing. Two-coder nominal alpha via coincidence matrix."""
    pr = [(x, y) for x, y in pairs if x is not None and y is not None]
    cats = sorted({v for p in pr for v in p})
    if not pr or len(cats) < 2:
        return float('nan')
    idx = {c: i for i, c in enumerate(cats)}
    o = np.zeros((len(cats), len(cats)))
    for x, y in pr:
        o[idx[x], idx[y]] += 1
        o[idx[y], idx[x]] += 1
    nc = o.sum(1); n = nc.sum()
    do = (o.sum() - np.trace(o))
    de = (nc.sum() ** 2 - (nc ** 2).sum())
    if de == 0:
        return float('nan')
    return 1.0 - (n - 1) * do / de


def cohen_kappa(pairs):
    pr = [(x, y) for x, y in pairs if x is not None and y is not None]
    n = len(pr)
    if n == 0:
        return float('nan')
    po = sum(x == y for x, y in pr) / n
    cats = {v for p in pr for v in p}
    pe = sum((sum(x == c for x, _ in pr) / n) * (sum(y == c for _, y in pr) / n) for c in cats)
    return float('nan') if pe == 1 else (po - pe) / (1 - pe)


def wilson(k, n):
    if n == 0:
        return (float('nan'), float('nan'))
    p = k / n; den = 1 + Z * Z / n
    c = (p + Z * Z / (2 * n)) / den
    h = Z * math.sqrt(p * (1 - p) / n + Z * Z / (4 * n * n)) / den
    return (c - h, c + h)


def binom_cdf(k, n, p):
    return sum(math.comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(0, k + 1))


def clopper_pearson(k, n, alpha=0.05):
    if n == 0:
        return (float('nan'), float('nan'))

    def bis(f, target_increasing):
        lo, hi = 0.0, 1.0
        for _ in range(100):
            mid = (lo + hi) / 2
            if (f(mid) > 0) == target_increasing:
                hi = mid
            else:
                lo = mid
        return (lo + hi) / 2
    # lower: P(X >= k | p) = alpha/2 ; P(X>=k) increasing in p
    lower = 0.0 if k == 0 else bis(lambda p: (1 - binom_cdf(k - 1, n, p)) - alpha / 2, True)
    # upper: P(X <= k | p) = alpha/2 ; decreasing in p
    upper = 1.0 if k == n else bis(lambda p: alpha / 2 - binom_cdf(k, n, p), True)
    return (lower, upper)


def boot_ci(fn, items, rng, B=2000):
    vals = []
    n = len(items)
    for _ in range(B):
        ix = rng.integers(0, n, n)
        v = fn([items[i] for i in ix])
        if not (isinstance(v, float) and math.isnan(v)):
            vals.append(v)
    if not vals:
        return (float('nan'), float('nan'), 0)
    return (float(np.percentile(vals, 2.5)), float(np.percentile(vals, 97.5)), len(vals))


def now():
    return datetime.datetime.now().isoformat(timespec='seconds')


# ================= STAGE 1 =================
if a.stage == 1:
    print('STAGE 1: opening key_calibration only')
    n_tot, miss_label, miss_lean, bad_conf = sheet_completeness()
    print(f'sheet check: rows={n_tot} missing_label={miss_label} U_without_lean={miss_lean} bad_confidence={bad_conf}')
    key = read_key('key_calibration.csv', ['passage_id', 'sample_id', 'cell', 'reference_label', 'nonhigh', 'e1_layer'])
    assert len(key) == 100, len(key)
    sh = read_sheet({k['passage_id'] for k in key})
    assert len(sh) == 100, 'calibration rows missing from the returned sheet'
    ref = []; c3 = []; lean = []
    for k in key:
        s = sh[k['passage_id']]
        assert s['label'] in ('1', '0', 'U'), f"bad label on {k['passage_id']}"
        ref.append(k['reference_label'])
        c3.append(s['label'])
        if s['label'] == 'U':
            assert s['lean'] in ('1', '0'), f"U without lean on {k['passage_id']}"
        lean.append(s['lean'])
    n_u = sum(x == 'U' for x in c3)
    pairs_miss = [(r, None if c == 'U' else c) for r, c in zip(ref, c3)]
    pairs_lean = [(r, l if c == 'U' else c) for r, c, l in zip(ref, c3, lean)]
    items = list(zip(ref, c3, lean))

    def a_miss(it):
        return kripp_alpha_nominal([(r, None if c == 'U' else c) for r, c, l in it])

    def a_lean(it):
        return kripp_alpha_nominal([(r, l if c == 'U' else c) for r, c, l in it])
    alpha_miss = a_miss(items); alpha_lean = a_lean(items)
    kappa = cohen_kappa(pairs_miss)
    nn = [(r, c) for r, c in pairs_miss if c is not None]
    agree = sum(r == c for r, c in nn)
    wl = wilson(agree, len(nn))
    with open(a.seed_file, encoding='utf-8') as f:
        seed = int(f.read().strip())
    rng = np.random.default_rng(seed)
    bm = boot_ci(a_miss, items, rng)
    bl = boot_ci(a_lean, items, rng)
    if n_u < 20:
        stat_name, stat = 'alpha_U_missing', alpha_miss
    elif n_u < 40:
        stat_name, stat = 'alpha_lean_substituted', alpha_lean
    else:
        stat_name, stat = 'alpha_lean_substituted', alpha_lean  # reported; outcome fixed by U >= 40
    qualified = bool((not math.isnan(stat)) and stat >= THRESH and n_u < 40)
    share_pos = sum(r == '1' for r in ref) / len(ref)
    nonhigh_share = sum(k['nonhigh'] == 'True' for k in key) / len(key)
    rec = dict(
        written_at=now(), sheet_sha256=sha(a.sheet), key_sha256=sha(os.path.join(a.keydir, 'key_calibration%s.csv' % suffix)),
        n_calibration=len(key), n_U=n_u, U_rate=n_u / len(key),
        alpha_U_missing=alpha_miss, alpha_U_missing_boot95=[bm[0], bm[1]], boot_valid_resamples_miss=bm[2],
        alpha_lean_substituted=alpha_lean, alpha_lean_boot95=[bl[0], bl[1]], boot_valid_resamples_lean=bl[2],
        cohen_kappa_U_excluded=kappa, percent_agreement_U_excluded=agree / len(nn) if nn else float('nan'),
        percent_agreement_n=len(nn), percent_agreement_wilson95=list(wl),
        qualification_statistic=stat_name, qualification_value=stat, threshold=THRESH,
        outcome='qualified' if qualified else 'not qualified',
        achieved_share_reference_positive=share_pos, nonhigh_share=nonhigh_share,
        bootstrap_seed_file_sha256=sha(a.seed_file), replacement=a.replacement)
    with open(REC, 'w', encoding='utf-8') as f:
        json.dump(rec, f, ensure_ascii=False, indent=2, sort_keys=True)
    h = sha(REC)
    with open(LOG, 'a', encoding='utf-8') as f:
        f.write(f'{now()}\t{os.path.basename(REC)}\t{h}\n')
    print('=== RESULTS ===')
    print(f'stage=1 replacement={a.replacement}')
    print(f'U={n_u}/100 alpha_U_missing={alpha_miss:.4f} CI[{bm[0]:.4f},{bm[1]:.4f}] alpha_lean={alpha_lean:.4f} CI[{bl[0]:.4f},{bl[1]:.4f}]')
    print(f'kappa_U_excluded={kappa:.4f} agreement={agree}/{len(nn)} Wilson[{wl[0]:.4f},{wl[1]:.4f}]')
    print(f'qualification_statistic={stat_name} value={stat:.4f} threshold={THRESH} outcome={rec["outcome"]}')
    print(f'achieved_share_reference_positive={share_pos:.3f} nonhigh_share={nonhigh_share:.3f}')
    print(f'qualification_record_sha256={h}')
    print('=== END RESULTS ===')
    sys.exit(0)

# ================= STAGE 2 =================
assert os.path.exists(REC) and os.path.exists(LOG), 'no qualification record: Stage 2 refused'
logged = [l.rstrip('\n').split('\t') for l in open(LOG, encoding='utf-8') if l.strip()]
cur = sha(REC)
assert logged and logged[-1][2] == cur, 'qualification record hash does not match the log: Stage 2 refused'
rec = json.load(open(REC, encoding='utf-8'))
qualified = rec['outcome'] == 'qualified'
print(f'STAGE 2: record hash verified ({cur[:12]}); outcome={rec["outcome"]}')

nc = read_key('key_nonconcordant.csv', ['passage_id', 'sample_id'])
e1 = read_key('key_e1.csv', ['passage_id', 'sample_id', 'stratum', 'reference_label'])
cal = read_key('key_calibration.csv', ['passage_id', 'sample_id', 'cell', 'reference_label', 'nonhigh', 'e1_layer'])
assert len(nc) == 123 and len(e1) == 214 and len(cal) == 100
sh = read_sheet()
assert len(sh) == 437, len(sh)
for r in nc + e1 + cal:
    s = sh[r['passage_id']]
    assert s['label'] in ('1', '0', 'U'), f"bad label {r['passage_id']}"
    if s['label'] == 'U':
        assert s['lean'] in ('1', '0'), f"U without lean {r['passage_id']}"

# concordant labels from the workbook (read-only)
wb = openpyxl.load_workbook(a.workbook, read_only=True, data_only=True)
it = wb['Annotation_Blind'].iter_rows(values_only=True)
h = list(next(it))
ix = {n: h.index(n) for n in ('sample_id', 'coder1_label', 'coder2_label')}
lab12 = {}
for r in it:
    if r[ix['sample_id']] is None:
        continue
    lab12[r[ix['sample_id']]] = (norm_label(r[ix['coder1_label']]), norm_label(r[ix['coder2_label']]))
assert len(lab12) == 600
conc = {s: v[0] for s, v in lab12.items() if v[0] == v[1] and v[0] in ('0', '1')}
assert len(conc) == 477

# --- non-concordant: final labels (7.2, 7.3) ---
final = dict(conc)
gold_lean = dict(conc); gold_u0 = dict(conc); gold_u1 = dict(conc)
nc_rows = []; unresolved = []
ag1 = ag2 = nbin = 0
for r in nc:
    s = sh[r['passage_id']]; sid = r['sample_id']
    c1, c2 = lab12[sid]; c3 = s['label']
    if c3 == 'U':
        unresolved.append(sid)
        gold_lean[sid] = s['lean']; gold_u0[sid] = '0'; gold_u1[sid] = '1'
        fin = ''
    else:
        final[sid] = c3; gold_lean[sid] = c3; gold_u0[sid] = c3; gold_u1[sid] = c3
        fin = c3
        nbin += 1
        ag1 += (c3 == c1); ag2 += (c3 == c2)
    nc_rows.append([sid, r['passage_id'], c1, c2, c3, s['lean'], s['confidence'], fin, s['reason']])
assert len(gold_lean) == 600

# --- calibration disagreements and flip variant (7.4) ---
cal_dis = []; flip = dict(gold_lean)
for r in cal:
    s = sh[r['passage_id']]
    if s['label'] in ('1', '0') and s['label'] != r['reference_label']:
        cal_dis.append([r['sample_id'], r['passage_id'], r['cell'], r['reference_label'], s['label'], s['confidence'], s['reason']])
        flip[r['sample_id']] = s['label']
    elif s['label'] == 'U':
        cal_dis.append([r['sample_id'], r['passage_id'], r['cell'], r['reference_label'], 'U(lean ' + s['lean'] + ')', s['confidence'], s['reason']])

# --- E1 (7.5, 8.2(ix)): 244 passages = 214 e1 + 30 NM/RND calibration ---
e1_all = []
for r in e1:
    e1_all.append((r['sample_id'], r['passage_id'], r['stratum'], sh[r['passage_id']]))
for r in cal:
    if r['e1_layer'] == 'True':
        e1_all.append((r['sample_id'], r['passage_id'], r['cell'].split('-')[0], sh[r['passage_id']]))
assert len(e1_all) == 244
e1_stats = {}
e1_ones = []; e1_us = []
for stratum in ('NM', 'RND', 'ALL'):
    sub = [x for x in e1_all if stratum == 'ALL' or x[2] == stratum]
    n = len(sub); k1 = sum(x[3]['label'] == '1' for x in sub); ku = sum(x[3]['label'] == 'U' for x in sub)
    lean1 = sum(x[3]['label'] == 'U' and x[3]['lean'] == '1' for x in sub)
    cp = clopper_pearson(k1, n)
    e1_stats[stratum] = dict(n=n, labeled_1=k1, CP95=list(cp), labeled_U=ku, U_lean1=lean1, U_lean0=ku - lean1)
for sid, pid, st, s in e1_all:
    if s['label'] == '1':
        e1_ones.append([sid, pid, st, s['confidence'], s['reason']])
    if s['label'] == 'U':
        e1_us.append([sid, pid, st, s['lean'], s['reason']])


def wcsv(name, header, rows):
    p = os.path.join(a.outdir, name)
    with open(p, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
    return p


wcsv('stage2_nonconcordant_labels.csv', ['sample_id', 'passage_id', 'coder1', 'coder2', 'coder3', 'coder3_lean', 'coder3_conf', 'final_label', 'coder3_reason'], nc_rows)
wcsv('stage2_calibration_disagreements.csv', ['sample_id', 'passage_id', 'cell', 'reference_label', 'coder3', 'coder3_conf', 'coder3_reason'], cal_dis)
wcsv('stage2_e1_labeled_1.csv', ['sample_id', 'passage_id', 'stratum', 'coder3_conf', 'coder3_reason'], e1_ones)
wcsv('stage2_e1_labeled_U.csv', ['sample_id', 'passage_id', 'stratum', 'lean', 'coder3_reason'], e1_us)
e1_p = os.path.join(a.outdir, 'stage2_e1_summary.json')
json.dump(e1_stats, open(e1_p, 'w', encoding='utf-8'), indent=2)

gold_written = False
if qualified:
    gp = os.path.join(a.outdir, 'final_labels.csv')
    rows = []
    for s in sorted(lab12):
        rows.append([s, final.get(s, ''), gold_lean[s], gold_u0[s], gold_u1[s], flip[s]])
    wcsv('final_labels.csv', ['sample_id', 'gold_primary', 'gold_lean', 'gold_U0', 'gold_U1', 'gold_flip_lean'], rows)
    gold_written = True

print('=== RESULTS ===')
print(f'stage=2 replacement={a.replacement} outcome={rec["outcome"]} record_sha256={cur}')
print(f'nonconcordant resolved(binary)={nbin} unresolved(U)={len(unresolved)} coder3_agrees_with_coder1={ag1}/{nbin} with_coder2={ag2}/{nbin}')
print(f'calibration disagreements/U listed={len(cal_dis)} flip variant label changes={sum(flip[s]!=gold_lean[s] for s in flip)}')
for k, v in e1_stats.items():
    print(f'E1 {k}: n={v["n"]} labeled_1={v["labeled_1"]} CP95=[{v["CP95"][0]:.4f},{v["CP95"][1]:.4f}] U={v["labeled_U"]} (lean1 {v["U_lean1"]}, lean0 {v["U_lean0"]})')
print(f'final_labels.csv written={gold_written} (written only if qualified; gold_primary blank where Coder 3 entered U)')
print('=== END RESULTS ===')
